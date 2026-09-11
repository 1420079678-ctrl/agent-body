#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify document conversion works with the magika shim (extension-based routing)."""
import sys, os, json

sys.stdout.reconfigure(encoding='utf-8')
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', 'lib', 'shims'))
sys.path.insert(0, os.path.join(HERE, '..', 'lib'))
sys.path.insert(0, os.path.join(HERE, '..', 'vendor'))

# 直接复用 worker 的转换函数，确认 shim 之后 PDF/DOCX/XLSX/CSV/HTML 都还能走对转换器。
import importlib.util
spec = importlib.util.spec_from_file_location('wcw', os.path.join(HERE, '..', 'lib', 'webcrawl_worker.py'))
wcw = importlib.util.module_from_spec(spec)
spec.loader.exec_module(wcw)

TESTS = [
    ('pdf',  'https://www.w3.org/WAI/ER/tests/xhtml/testfiles/resources/pdf/dummy.pdf'),
    ('txt',  'https://raw.githubusercontent.com/mwaskom/seaborn-data/master/README.md'),
    ('html', 'https://example.com/'),
    ('csv',  'https://raw.githubusercontent.com/mwaskom/seaborn-data/master/iris.csv'),
    ('xlsx', 'https://go.microsoft.com/fwlink/?LinkID=521962'),
]

for name, url in TESTS:
    r = wcw.doc_to_markdown(url=url, max_length=4000)
    ok = r.get('success')
    print(f"[{name:5}] ok={ok} ext={r.get('ext')} chars={r.get('chars')} "
          f"head={(r.get('markdown') or r.get('error') or '')[:90]!r}")

# 本地 DOCX（用 office-docs 生成一个再转）
try:
    from docx import Document
    p = os.path.join(os.environ.get('TEMP', '.'), 'wcw_probe.docx')
    d = Document()
    d.add_heading('Shim check', level=1)
    d.add_paragraph('magika shim must not break docx routing.')
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text = 'a'; t.cell(0, 1).text = 'b'
    d.save(p)
    r = wcw.doc_to_markdown(path=p, max_length=2000)
    print(f"[docx ] ok={r.get('success')} ext={r.get('ext')} chars={r.get('chars')} "
          f"head={(r.get('markdown') or r.get('error') or '')[:90]!r}")
except Exception as e:
    print(f'[docx ] skipped: {type(e).__name__}: {e}')
