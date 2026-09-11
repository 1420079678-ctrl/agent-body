#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dsh-office-docs 引擎：源自 Sol 5.6 skills（PDF/DOCX/PPTX/XLSX）的确定性执行层。
用法: python office_docs.py <action>，JSON 参数从 stdin 读取（绕开 Windows argv 编码问题）。
stdout 输出 JSON 结果（UTF-8）；任何异常写 stderr 并以 exit 1 退出。
"""
import json, os, sys, tempfile, subprocess, shutil

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")
if hasattr(sys.stdin, "reconfigure"):
    sys.stdin.reconfigure(encoding="utf-8")  # 关键：Windows 管道下 stdin 默认 GBK，必须强制 UTF-8

def out(obj):
    print(json.dumps(obj, ensure_ascii=False, default=str))

def err(msg):
    sys.stderr.write(str(msg) + "\n")
    sys.exit(1)

def load_args(raw):
    try:
        return json.loads(raw) if raw else {}
    except Exception as e:
        err(f"args JSON 解析失败: {e}")

# ─────────────────────────── PDF（reportlab + pdfplumber + pypdf）───────────────────────────

def pdf_build(a):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    # 注册中文字体（Windows 自带）。优先单 .ttf（TTC 子字体在 reportlab 中 ToUnicode 映射有缺陷，
    # 会导致 PDF 提取乱码）；.ttf 缺失才回退 TTC。
    CJK_FONTS = [
        ("CJK-simhei", r"C:\Windows\Fonts\simhei.ttf", None),   # 黑体（单 ttf，首选）
        ("CJK-simkai", r"C:\Windows\Fonts\simkai.ttf", None),   # 楷体
        ("CJK-simfang", r"C:\Windows\Fonts\simfang.ttf", None),  # 仿宋
        ("CJK-deng", r"C:\Windows\Fonts\Deng.ttf", None),       # 等线
        ("CJK-yahei", r"C:\Windows\Fonts\msyh.ttc", 0),          # 微软雅黑（TTC 兜底）
        ("CJK-simsun", r"C:\Windows\Fonts\simsun.ttc", 0),       # 宋体（TTC 兜底）
    ]
    registered = None
    for name, fp, sub in CJK_FONTS:
        try:
            if os.path.exists(fp):
                if sub is not None and fp.lower().endswith(".ttc"):
                    pdfmetrics.registerFont(TTFont(name, fp, subfontIndex=sub))
                else:
                    pdfmetrics.registerFont(TTFont(name, fp))
                registered = name
                break
        except Exception:
            continue
    if not registered:
        err("未找到可用中文字体（simhei.ttf/msyh.ttc/simsun.ttc），无法正确渲染中文 PDF")

    path = a.get("path") or os.path.join(os.getcwd(), "output.pdf")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc = SimpleDocTemplate(path, pagesize=A4, leftMargin=20*mm, rightMargin=20*mm, topMargin=18*mm, bottomMargin=18*mm)
    styles = getSampleStyleSheet()
    styles["Title"].fontName = registered
    styles["Heading1"].fontName = registered
    styles["Heading2"].fontName = registered
    styles["BodyText"].fontName = registered
    styles["BodyText"].fontSize = 11
    story = []
    for blk in a.get("blocks", []):
        t = blk.get("type", "para")
        text = blk.get("text", "")
        if t == "heading":
            lvl = int(blk.get("level", 1))
            style = styles["Title"] if lvl == 0 else (styles["Heading1"] if lvl == 1 else styles["Heading2"])
            story.append(Paragraph(text, style))
        elif t == "para":
            style = styles["BodyText"] if not blk.get("bold") else styles["BodyText"]
            story.append(Paragraph(text, style))
        elif t == "table":
            rows = blk.get("rows", [[""]])
            tstyle = ParagraphStyle("cell", fontName=registered, fontSize=9.5, leading=12)
            cell_rows = [[Paragraph(str(c), tstyle) for c in row] for row in rows]
            tb = Table(cell_rows)
            tb.setStyle(TableStyle([
                ("GRID", (0,0), (-1,-1), 0.5, colors.grey),
                ("BACKGROUND", (0,0), (-1,0), colors.lightgrey),
                ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
            ]))
            story.append(tb)
        elif t == "image":
            img = blk.get("path") or blk.get("src")
            if img and os.path.exists(img):
                w = blk.get("width", 120)
                story.append(Image(img, width=w*mm, height=w*mm*(blk.get("ratio") or 0.75)))
        elif t == "spacer":
            story.append(Spacer(1, int(blk.get("height", 8))))
        story.append(Spacer(1, 4))
    doc.build(story)
    return {"ok": True, "path": os.path.abspath(path), "pages": 1, "font": registered, "note": "reportlab 构建完成，建议用 pdf_render 渲染 PNG 做视觉验证"}

def pdf_extract(a):
    path = a["path"]
    if not os.path.exists(path):
        err(f"文件不存在: {path}")
    result = {"path": os.path.abspath(path), "pages": 0, "text": "", "tables": [], "meta": {}}
    try:
        with open(path, "rb") as f:
            from pypdf import PdfReader
            r = PdfReader(f)
            result["pages"] = len(r.pages)
            result["meta"] = {k: str(v) for k, v in (r.metadata or {}).items()}
    except Exception as e:
        result["meta"]["note"] = f"pypdf 读取失败: {e}"
    try:
        with pdfplumber_open(path) as pdf:
            result["pages"] = len(pdf.pages)
            texts = []
            for p in pdf.pages:
                texts.append(p.extract_text() or "")
                for tb in p.extract_tables() or []:
                    if tb:
                        result["tables"].append([[c or "" for c in row] for row in tb])
            result["text"] = "\n\n".join(texts)
    except Exception as e:
        result.setdefault("meta", {})["note2"] = f"pdfplumber 失败: {e}"
    return result

def pdfplumber_open(path):
    import pdfplumber
    return pdfplumber.open(path)

def pdf_render(a):
    """用 pypdfium2 渲染 PDF 每页为 PNG（Sol 5.6 视觉验证替代 poppler/pdftoppm）。"""
    import pypdfium2 as pdfium
    path = a["path"]
    outdir = a.get("outdir") or os.path.join(os.path.dirname(os.path.abspath(path)), "render")
    scale = float(a.get("scale", 1.5))
    os.makedirs(outdir, exist_ok=True)
    pdf = pdfium.PdfDocument(path)
    pngs = []
    for i, page in enumerate(pdf):
        bitmap = page.render(scale=scale)
        png = os.path.join(outdir, f"page-{i+1:03d}.png")
        bitmap.to_pil().save(png)
        pngs.append(os.path.abspath(png))
    return {"ok": True, "pages": len(pngs), "pngs": pngs, "outdir": os.path.abspath(outdir)}

# ─────────────────────────── DOCX（python-docx）───────────────────────────

def docx_build(a):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    path = a.get("path") or os.path.join(os.getcwd(), "output.docx")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    doc = Document()
    for blk in a.get("blocks", []):
        t = blk.get("type", "para")
        text = blk.get("text", "")
        if t == "heading":
            doc.add_heading(text, level=int(blk.get("level", 1)))
        elif t == "para":
            p = doc.add_paragraph(text)
            if blk.get("bold"):
                for run in p.runs:
                    run.bold = True
        elif t == "table":
            rows = blk.get("rows", [[""]])
            tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
            tb.style = "Light Grid Accent 1"
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    tb.cell(i, j).text = str(cell)
        elif t == "pagebreak":
            doc.add_page_break()
    doc.save(path)
    return {"ok": True, "path": os.path.abspath(path)}

def docx_extract(a):
    from docx import Document
    path = a["path"]
    if not os.path.exists(path):
        err(f"文件不存在: {path}")
    doc = Document(path)
    paras = [p.text for p in doc.paragraphs]
    tables = [[[c.text for c in row.cells] for row in tb.rows] for tb in doc.tables]
    return {
        "path": os.path.abspath(path),
        "paragraphs": len(paras),
        "tables": len(tables),
        "text": "\n".join(paras),
        "tables_data": tables,
    }

def docx_append(a):
    from docx import Document
    path = a["path"]
    if not os.path.exists(path):
        err(f"文件不存在: {path}")
    doc = Document(path)
    for blk in a.get("blocks", []):
        t = blk.get("type", "para")
        text = blk.get("text", "")
        if t == "heading":
            doc.add_heading(text, level=int(blk.get("level", 1)))
        elif t == "table":
            rows = blk.get("rows", [[""]])
            tb = doc.add_table(rows=len(rows), cols=len(rows[0]))
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    tb.cell(i, j).text = str(cell)
        else:
            doc.add_paragraph(text)
    doc.save(path)
    return {"ok": True, "path": os.path.abspath(path)}

# ─────────────────────────── PPTX（python-pptx）───────────────────────────

def pptx_build(a):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    path = a.get("path") or os.path.join(os.getcwd(), "output.pptx")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    prs = Presentation()
    for s in a.get("slides", []):
        layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = s.get("title", "")
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = ""
        for i, bullet in enumerate(s.get("bullets", [])):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = bullet
            p.level = 0
    prs.save(path)
    return {"ok": True, "path": os.path.abspath(path), "slides": len(a.get("slides", []))}

def pptx_extract(a):
    from pptx import Presentation
    path = a["path"]
    if not os.path.exists(path):
        err(f"文件不存在: {path}")
    prs = Presentation(path)
    slides = []
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(shape.text)
        slides.append({"index": i + 1, "text": "\n".join(texts)})
    return {"path": os.path.abspath(path), "slides": len(slides), "slides_data": slides}

# ─────────────────────────── XLSX（openpyxl）───────────────────────────

def xlsx_build(a):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    path = a.get("path") or os.path.join(os.getcwd(), "output.xlsx")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    for sh in a.get("sheets", []):
        ws = wb.create_sheet(sh.get("name", "Sheet1"))
        for r, row in enumerate(sh.get("rows", []), start=1):
            for c, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c)
                cell.value = val
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="DDDDDD")
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(path)
    return {"ok": True, "path": os.path.abspath(path), "sheets": wb.sheetnames}

def xlsx_extract(a):
    from openpyxl import load_workbook
    path = a["path"]
    if not os.path.exists(path):
        err(f"文件不存在: {path}")
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        sheets.append({"name": ws.title, "rows": rows})
    return {"path": os.path.abspath(path), "sheets": len(sheets), "sheets_data": sheets}

# ─────────────────────────── 通用转换（pandoc）───────────────────────────

def convert(a):
    pandoc = shutil.which("pandoc")
    if not pandoc:
        err("pandoc 未安装：无法做格式转换（可用 docx_build/pptx_build/xlsx_build 直接构建）")
    src, dst = a.get("input"), a.get("output")
    if not src or not dst:
        err("需要 input 与 output 路径")
    if not os.path.exists(src):
        err(f"输入文件不存在: {src}")
    FMT_ALIAS = {"md": "markdown", "mdown": "markdown", "mkd": "markdown", "htm": "html", "rtf": "rtf"}
    raw_from = a.get("from") or os.path.splitext(src)[1].lstrip(".").lower()
    raw_to = a.get("to") or os.path.splitext(dst)[1].lstrip(".").lower()
    from_fmt = FMT_ALIAS.get(raw_from, raw_from)
    to_fmt = FMT_ALIAS.get(raw_to, raw_to)
    cmd = [pandoc, src, "-f", from_fmt, "-t", to_fmt, "-o", dst]
    if a.get("standalone"):
        cmd.append("-s")
    os.makedirs(os.path.dirname(os.path.abspath(dst)), exist_ok=True)
    r = subprocess.run(cmd, capture_output=True, text=True)
    if r.returncode != 0:
        err(f"pandoc 失败: {r.stderr[:500]}")
    return {"ok": True, "input": os.path.abspath(src), "output": os.path.abspath(dst), "from": from_fmt, "to": to_fmt}

# ─────────────────────────── 路由 ───────────────────────────

ACTIONS = {
    "pdf_build": pdf_build,
    "pdf_extract": pdf_extract,
    "pdf_render": pdf_render,
    "docx_build": docx_build,
    "docx_extract": docx_extract,
    "docx_append": docx_append,
    "pptx_build": pptx_build,
    "pptx_extract": pptx_extract,
    "xlsx_build": xlsx_build,
    "xlsx_extract": xlsx_extract,
    "convert": convert,
}

def main():
    if len(sys.argv) < 2:
        err("用法: python office_docs.py <action>，JSON 参数从 stdin 读取")
    action = sys.argv[1]
    if action not in ACTIONS:
        err(f"未知 action: {action}（可用: {', '.join(ACTIONS)}）")
    raw = sys.stdin.read() if not sys.stdin.isatty() else None
    a = load_args(raw)
    try:
        out(ACTIONS[action](a))
    except Exception as e:
        import traceback
        err(f"{action} 执行失败: {e}\n{traceback.format_exc()}")

if __name__ == "__main__":
    main()
